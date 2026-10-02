# -*- coding: utf-8 -*-
"""Generate HOMEJI_WEB_SYSTEM.docx — agent memory for Homeji web + backend."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Pt, RGBColor

OUT = Path(__file__).resolve().parent / "HOMEJI_WEB_SYSTEM.docx"


def set_run(run, *, bold=False, size=11, color=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for r in h.runs:
        r.font.name = "Calibri"
    return h


def add_p(doc, text, *, bold=False, size=11):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(text)
    set_run(run, bold=bold, size=size)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(item)
        set_run(run, size=11)


def add_flow(doc, title, goal, route, apis, states):
    add_heading(doc, title, 3)
    add_p(doc, f"Mục tiêu người dùng: {goal}")
    add_p(doc, f"Route / màn hình web: {route}")
    add_p(doc, "API chính:")
    add_bullets(doc, apis)
    add_p(doc, "Trạng thái quan trọng:")
    add_bullets(doc, states)


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Pt(56)
    section.bottom_margin = Pt(56)
    section.left_margin = Pt(64)
    section.right_margin = Pt(64)

    add_heading(doc, "Homeji Web — Bộ nhớ hệ thống cho agent", 0)
    add_p(
        doc,
        "Tài liệu này tổng hợp kiến trúc và luồng thật từ mã nguồn (đã đối chiếu code). "
        "Mục đích: agent sau đọc nhanh để không phá route/API khi refactor UI. "
        "Không chứa secret / connection string / giá trị .env.",
    )
    add_p(
        doc,
        "Phạm vi: exe-homeji-web + exe-homeji-backend. Không mô tả exe-homeji-mobile "
        "(app độc lập, chỉ dùng chung backend path).",
        bold=True,
    )
    add_p(
        doc,
        "Ngày lập: 2026-09-30. Nguồn chính: PRODUCT.md, DESIGN.md, "
        "exe-homeji-web/src/App.tsx, api/*, components/map/*, "
        "exe-homeji-backend/Homeji_BE/src/Homeji.Api/.",
    )

    # ── A. Product & roles ──
    add_heading(doc, "A. Sản phẩm và vai trò", 1)
    add_p(
        doc,
        "Homeji là nền tảng tìm phòng trọ và bạn ở ghép (khu vực Thủ Đức & Quận 9). "
        "Slogan landing: 「Trọ an tâm · Nâng tầm cuộc sống」. "
        "Web là SPA React; mobile là app riêng — không import/copy source giữa hai app.",
    )
    add_heading(doc, "Vai trò trong code (UserRole)", 2)
    add_bullets(
        doc,
        [
            "Guest (chưa đăng nhập): không có enum — chỉ là trạng thái !isAuthenticated. "
            "Landing cinematic tại `/` (HomePage khi chưa auth), bản đồ công khai trong GuestMapSection, "
            "auth tại `/login` `/register`.",
            "Renter = 1 — người thuê / tìm phòng (primary). Có thể đăng tin pass phòng (`/posts/new?type=pass`). "
            "File: exe-homeji-web/src/api/types.ts; backend: Homeji.Domain/Enums/UserRole.cs.",
            "Landlord = 2 — chủ nhà / người đăng tin thuê. Đăng tin `/posts/new`, quản lý Tin của tôi, "
            "xác minh chủ nhà (landlord-verifications).",
            "Admin = 3 — kiểm duyệt tại `/admin` (AdminRoute kiểm tra profile.role === UserRole.Admin).",
        ],
    )
    add_p(
        doc,
        "Lưu ý: UI marketing nói “sinh viên / chủ nhà”; enum kỹ thuật chỉ có Renter | Landlord | Admin. "
        "Profile có thể thiếu ngay sau đăng ký (404 /api/profile/me → needsProfileSetup).",
    )

    # ── B. Backend ──
    add_heading(doc, "B. Backend và cách web gọi API", 1)
    add_heading(doc, "Vị trí backend", 2)
    add_bullets(
        doc,
        [
            "Thư mục: exe-homeji-backend/Homeji_BE/ (solution Homeji.sln).",
            "API host: Homeji.Api (ASP.NET Core) — Controllers + SignalR hub.",
            "Layer: Homeji.Domain / Application / Infrastructure / Api (+ tests).",
            "Deploy mặc định mà Vite proxy trỏ tới (không phải secret): biến HOMEJI_API_URL "
            "hoặc fallback https://homeji-be.onrender.com — xem exe-homeji-web/vite.config.ts.",
        ],
    )
    add_heading(doc, "Cách web nối tới backend", 2)
    add_bullets(
        doc,
        [
            "Base URL: resolveApiBase() trong exe-homeji-web/src/api/client.ts — Dev: same-origin '' "
            "(Vite proxy `/api` và `/hubs`). Prod: '' (proxy serve) hoặc VITE_API_BASE_URL nếu set.",
            "Auth HTTP: header Authorization: Bearer <accessToken>. Token lưu localStorage "
            "key homeji_access_token (+ homeji_user_id, homeji_email). Cơ chế: exe-homeji-web/src/api/authSession.ts.",
            "Không dùng cookie session cho API JSON thông thường.",
            "Backend xác thực JWT Supabase (JWKS, audience/issuer từ config) — "
            "Homeji.Api/Authentication/SupabaseAuthenticationExtensions.cs. "
            "Hub SignalR nhận access_token qua query string khi path /hubs.",
            "Realtime: web hook useNotificationHub → /hubs/notifications "
            "(exe-homeji-web/src/hooks/useNotificationHub.ts); backend MapHub NotificationHub "
            "trong Program.cs.",
        ],
    )
    add_heading(doc, "Nhóm API chính (Route controllers)", 2)
    add_bullets(
        doc,
        [
            "api/account — đăng ký, login, quên/đặt lại MK, Google OAuth URL/redirect, email-availability.",
            "api/profile — /me, lifestyle.",
            "api/rental-posts — search, detail, draft/update/media/submit/archive/mark-rented, compare, mine/stats; "
            "nested viewing-appointments create; reviews.",
            "api/saved-posts — lưu/bỏ lưu, roommate-candidates.",
            "api/roommate-invitations — mine, create, accept/reject/cancel.",
            "api/conversations — inbox, start theo rental/marketplace/wanted, messages + ảnh.",
            "api/viewing-appointments — list + confirm/reject/cancel/reschedule/complete.",
            "api/marketplace-posts + api/marketplace-orders — chợ đồ + đơn/cart.",
            "api/rental-wanted-posts — tin tìm phòng.",
            "api/notifications — list, read, read-all.",
            "api/activities — nhật ký hoạt động.",
            "api/subscriptions + api/payments + api/wallet — gói, thanh toán MoMo/PayOS, ví/rút tiền.",
            "api/ai + api/chatbot — parse search, highlight, chatbot Homie.",
            "api/upload — ảnh.",
            "api/reports — báo cáo nội dung.",
            "api/landlord-verifications — xác minh chủ nhà.",
            "api/admin/moderation + api/admin/wallet-withdrawals + api/admin/landlord-verifications — admin.",
            "api/roommate-chats — (controller riêng; web chủ yếu dùng conversations).",
        ],
    )
    add_p(
        doc,
        "Client wrapper đầy đủ: exe-homeji-web/src/api/index.ts (export hàm gọi từng path).",
    )

    add_heading(doc, "B2. Luồng end-to-end (đã đối chiếu code)", 1)

    add_flow(
        doc,
        "1) Login / session",
        "Đăng nhập email/password hoặc Google; giữ phiên; load hồ sơ.",
        "/login, /register, /forgot-password, /reset-password, /auth/callback; "
        "AuthContext (contexts/AuthContext.tsx); modal AuthModalContext.",
        [
            "POST /api/account/login | register → AuthSession.accessToken",
            "GET /api/account/google/url (+ redirect); callback đưa token vào session",
            "GET /api/profile/me — 404 → needsProfileSetup; 401 → clear session",
        ],
        [
            "hasToken / isLoading / profile / needsProfileSetup / authDisrupted",
            "Token hết hạn: expireStoredAuth + event homeji:auth-expired",
            "ProtectedRoute chặn khi chưa auth; AdminRoute yêu cầu role Admin",
        ],
    )

    add_flow(
        doc,
        "2) Listings / search (Tìm phòng)",
        "Tìm/lọc tin thuê quanh Thủ Đức–Q.9, xem danh sách + ghim.",
        "Authed: `/` + section listings (panel). Guest: landing #map / GuestMapSection; /explore → /#map.",
        [
            "GET /api/rental-posts (searchRentalPosts) — filter giá, loại, amenities…",
            "POST /api/ai/parse-search, POST /api/ai/highlight-rental-posts (omnibox/AI)",
        ],
        [
            "posts[], selectedPostId, panelSection='listings', pin layers (mapPinLayers)",
            "HomePage loadPosts khi authenticated; MapOmnibox filter/chips",
        ],
    )

    add_flow(
        doc,
        "3) Map pins & place detail",
        "Click ghim / danh sách → focus map → mở chi tiết chỗ ở; lưu / nhắn / hẹn từ panel.",
        "`/?post=<id>` deep link (MapHomePostRedirect từ /posts/:postId); "
        "MapPlaceDetailPanel trong AuthenticatedHomeMapShell.",
        [
            "GET /api/rental-posts/{id}",
            "PUT/DELETE /api/saved-posts/{id}",
            "POST /api/conversations/rental-posts/{id}",
            "POST /api/rental-posts/{id}/viewing-appointments",
            "Reviews: GET/PUT reviews…",
        ],
        [
            "selectedPost / selectedPlace; panel detail overlay trên map",
            "Renter: lưu + đặt lịch; Landlord/owner: quản lý trạng thái tin",
            "Không có trang detail độc lập — redirect về map home",
        ],
    )

    add_flow(
        doc,
        "4) Saved (Đã lưu)",
        "Xem lại tin đã bookmark.",
        "`/?section=saved` hoặc redirect /saved → MapAppPanel → SavedPostsPage embedded.",
        ["GET /api/saved-posts", "DELETE /api/saved-posts/{id}", "GET …/roommate-candidates"],
        ["panelSection saved; WIDE_MAP_SECTIONS"],
    )

    add_flow(
        doc,
        "5) Marketplace (Chợ đồ)",
        "Mua/bán đồ cũ quanh khu trọ; pin marketplace trên map; đơn hàng/cart.",
        "`/?section=marketplace` / redirect /marketplace → MarketplacePage embedded.",
        [
            "GET/POST /api/marketplace-posts…",
            "POST orders, cart /api/marketplace-orders…",
            "POST /api/conversations/marketplace-posts/{id}",
            "Wallet khi liên quan thanh toán seller",
        ],
        ["marketplacePins trên RentalMap; cart open change; selectedMarketplaceId"],
    )

    add_flow(
        doc,
        "6) Wanted (Tin tìm phòng)",
        "Đăng nhu cầu thuê để chủ/người khác liên hệ.",
        "`/?section=wanted` / /wanted → WantedPostsPage embedded.",
        [
            "GET/POST/PUT /api/rental-wanted-posts…",
            "POST …/close",
            "POST /api/conversations/rental-wanted-posts/{id}",
        ],
        ["panel wanted; slide-from-right giống listings"],
    )

    add_flow(
        doc,
        "7) Invitations / roommate (Ở ghép)",
        "Gửi/nhận lời mời ở ghép; mở chat liên quan.",
        "`/?section=invitations` / /invitations → RoommateInvitationsPage.",
        [
            "GET /api/roommate-invitations/mine",
            "POST create/accept/reject/cancel",
            "Candidates từ saved-posts",
        ],
        ["panel invitations; onOpenConversation → MapChatDock"],
    )

    add_flow(
        doc,
        "8) Messages (Tin nhắn)",
        "Inbox hội thoại theo tin đăng.",
        "Không có route riêng; mở qua omnibox/nav-rail section messages → MapChatDock "
        "(AuthenticatedHomeMapShell: panelSection==='messages' không render MapAppPanel content).",
        [
            "GET /api/conversations",
            "GET/POST …/messages (+ images upload)",
        ],
        [
            "Chat dock state riêng; deep-link HomePage allowed sections hiện KHÔNG gồm 'messages' "
            "(/?section=messages bị bỏ qua trong effect deep link — chỉ mở bằng UI).",
        ],
    )

    add_flow(
        doc,
        "9) Appointments (Lịch xem phòng)",
        "Tạo/xác nhận/đổi lịch xem phòng.",
        "`/?section=appointments` → MapAppointmentsPanel; CTA từ MapPlaceDetailPanel.",
        [
            "GET /api/viewing-appointments",
            "POST create trên rental-posts/…/viewing-appointments",
            "confirm / reject / cancel / reschedule / complete",
        ],
        ["panel appointments (wide); notification deep-open"],
    )

    add_flow(
        doc,
        "10) Payments / subscriptions (Gói Đăng ký)",
        "Xem gói, thanh toán MoMo/PayOS, theo dõi payment return URL.",
        "`/?section=payments` / /payments (+ query paymentId/orderCode được giữ bởi mapSectionDeepLink).",
        [
            "GET /api/subscriptions/packages, /me",
            "POST premium …/momo|payos/create",
            "GET /api/payments, /orders/{orderCode}",
            "Wallet: /api/wallet…",
        ],
        [
            "PaymentPage embedded; return gateway giữ query trên /?section=payments&…",
            "MoMo orderId được normalize → orderCode trong mapDeepLinks.ts",
        ],
    )

    add_flow(
        doc,
        "11) Notifications + Activities",
        "Thông báo realtime + nhật ký thao tác.",
        "`/?section=notifications`, `/?section=activities`.",
        [
            "GET/POST /api/notifications… (+ default client-side notifications trong lib/defaultNotifications)",
            "SignalR notificationReceived",
            "GET /api/activities",
        ],
        ["badge unread trên omnibox; notification → open section liên quan"],
    )

    add_flow(
        doc,
        "12) Profile + My posts + Đăng tin",
        "Hồ sơ/lifestyle; quản lý tin; form tạo/sửa (trang thật).",
        "Profile/MyPosts: sections trên map. Tạo/sửa: `/posts/new`, `/posts/:id/edit` (ProtectedRoute) — "
        "KHÔNG embed trong map panel.",
        [
            "GET/PUT /api/profile/me, PUT …/lifestyle",
            "Draft/update/submit/archive rental-posts; upload /api/upload/image",
            "GET /api/rental-posts/mine/stats; landlord-verifications",
        ],
        ["needsProfileSetup banner; role quyết định CTA Đăng tin"],
    )

    add_flow(
        doc,
        "13) Admin",
        "Duyệt tin, báo cáo, user session, rút ví, xác minh landlord, thông báo bảo trì.",
        "Route thật `/admin` — AdminModerationPage (không phải map section).",
        [
            "api/admin/moderation/*",
            "api/admin/wallet-withdrawals/*",
            "api/admin/landlord-verifications/*",
        ],
        ["Chỉ Admin; terminate session user → client clear token"],
    )

    add_p(
        doc,
        "Ghi chú 'So sánh': API POST /api/rental-posts/compare tồn tại trong api/index.ts, "
        "nhưng không có màn/nav chính tên 'So sánh'. Không coi là feature UI cấp 1 khi redesign.",
        bold=True,
    )

    # ── C. Frontend map ──
    add_heading(doc, "C. Bản đồ frontend: router, shell, sections", 1)
    add_heading(doc, "Router (App.tsx)", 2)
    add_bullets(
        doc,
        [
            "Entry: src/main.tsx → App.tsx (BrowserRouter, AuthProvider, GoogleMapsProvider, …).",
            "Layout: AppLayout — Navbar + Outlet + footer + MobileTabBar (có điều kiện ẩn).",
            "Trang thật: /, /explore, auth pages, /posts/new, /posts/:id/edit, /admin.",
            "Redirect section: /my-posts, /marketplace, /wanted, /activities, /saved, /profile, "
            "/notifications, /invitations, /payments → MapHomeSectionRedirect → /?section=…",
            "/posts/:postId → MapHomePostRedirect → /?post=…",
            "* → Navigate /",
        ],
    )
    add_heading(doc, "Layout shells & ẩn Navbar", 2)
    add_bullets(
        doc,
        [
            "AppLayout: isMapHome = authenticated && (pathname==='/' || isMapSectionRedirectPath). "
            "Khi isMapHome: main-map-layout, không MobileTabBar; CSS app-shell--map đặt --nav-height: 0.",
            "Navbar.tsx: return null khi guest landing HOẶC map home (authed `/`). "
            "Trên map, điều hướng thật nằm trong MapOmnibox (nav-rail + hamburger drawer) + MapAccountMenu.",
            "Guest landing: chrome GuestChrome; map nhúng GuestMapSection.",
        ],
    )
    add_heading(doc, "Section ids (MapAppSection)", 2)
    add_p(
        doc,
        "Định nghĩa: exe-homeji-web/src/components/map/MapAppPanel.tsx — "
        "listings | saved | invitations | notifications | messages | appointments | payments | "
        "profile | marketplace | wanted | activities | myPosts.",
    )
    add_p(
        doc,
        "Nhãn UI drawer (MapOmnibox): Tìm phòng, Chợ đồ, Tin tìm phòng, Đã lưu, Ở ghép, "
        "Thông báo, Tin nhắn, Lịch xem phòng, Gói Đăng Ký, Nhật ký hoạt động, Tin của tôi, Đăng tin, Quản trị.",
    )
    add_heading(doc, "Ownership component", 2)
    add_bullets(
        doc,
        [
            "HomePage.tsx — phân nhánh guest landing vs authed; state posts/selection/section/searchParams.",
            "AuthenticatedHomeMapShell.tsx — khung map full-viewport: RentalMap + omnibox + panels + chat + chatbot.",
            "MapOmnibox.tsx — tìm kiếm Places, filter pin, nav-rail (Đã lưu/Chat/Lịch/Ở ghép), drawer menu, tour.",
            "MapAppPanel.tsx — host các page embedded (SavedPostsPage, MarketplacePage, …).",
            "MapPlaceDetailPanel.tsx — chi tiết tin/place; save/schedule/chat CTAs.",
            "MapChatDock.tsx — inbox/messages (thay panel khi section=messages).",
            "MapChatbot.tsx — Homie; có thể mở section.",
            "RentalMap.tsx — Google Maps Vector/WebGL, pins, zoom, view modes.",
            "Pages *Page.tsx còn hỗ trợ prop embedded để chạy trong panel hoặc (hiếm) standalone.",
        ],
    )
    add_heading(doc, "Deep link hành vi", 2)
    add_bullets(
        doc,
        [
            "HomePage đọc ?section= và ?post= rồi xóa khỏi URL (replace) sau khi apply state.",
            "allowed deep-link sections (HomePage) thiếu 'messages' — cần nhớ khi giữ URL tương thích.",
            "mapSectionDeepLink giữ query thanh toán khi redirect từ /payments.",
        ],
    )

    # ── D. Must not break ──
    add_heading(doc, "D. Không được phá khi viết lại UI", 1)
    add_bullets(
        doc,
        [
            "Routes trong App.tsx: giữ path cũ hoạt động (redirect hoặc trang). "
            "Không đổi tên/xóa route mà không có quyết định user + redirect tương thích.",
            "Section id strings trong MapAppSection + MAP_SECTION_REDIRECT_PATHS + mapSectionUrl.",
            "Query ?section= / ?post= / payment return (paymentId, orderCode, orderId→orderCode).",
            "API path + shape request/response (api/index.ts + types.ts). Không đổi contract backend trong UI rewrite.",
            "Auth: Bearer token localStorage + AuthContext semantics (needsProfileSetup, AdminRoute).",
            "SignalR /hubs/notifications URL và event notificationReceived.",
            "Map data logic: fetch pins, fit bounds, Google Maps provider — không đụng khi chỉ đổi chrome/nav.",
            "Form đăng tin routes /posts/new và /posts/:id/edit vẫn là trang riêng (ProtectedRoute).",
            "Admin /admin vẫn là trang riêng.",
            "Không import từ exe-homeji-mobile; không thêm Next/Tailwind/shadcn trừ khi user duyệt dependency.",
            "Ràng buộc map GPU: không StrictMode double-mount Vector; tránh View Transitions trên map home; "
            "--map-canvas-bg không đổi sang brand green.",
        ],
    )

    # ── E. Agent memory ──
    add_heading(doc, "E. Agent memory (ngắn)", 1)
    add_heading(doc, "Stack hiện tại", 2)
    add_bullets(
        doc,
        [
            "React 19 + TypeScript + Vite + react-router-dom + plain CSS / CSS variables (src/index.css).",
            "Không Next.js, không RSC, không Tailwind, không shadcn (stack lock).",
            "Backend: ASP.NET Core + Supabase JWT + SignalR.",
            "Docs sản phẩm: PRODUCT.md, DESIGN.md (root EXE101 hardlink tới exe-homeji-web).",
        ],
    )
    add_heading(doc, "Design tokens hôm nay", 2)
    add_bullets(
        doc,
        [
            "Brand: --grab-green #00b14f, --grab-green-hover #009241, --grab-blue #136fd8.",
            "Maps chrome: --maps-blue #1a73e8, soft/text variants; --map-canvas-bg #e8eef2.",
            "Surface ladder: --surface, --surface-muted, --surface-input, --border, text-* .",
            "Font app: Inter (+ Roboto/system). Landing: Manrope + Be Vietnam Pro. Auth: Sora/Manrope.",
            "Utility classes: .btn, .card, .form-*, .container, .page-* . Dark: html[data-theme] + prefers.",
            "Map home: --nav-height: 0 trong app-shell--map.",
        ],
    )
    add_heading(doc, "Hướng mới từ user (override PRODUCT 'map là ngôi nhà')", 2)
    add_bullets(
        doc,
        [
            "Refactor TOÀN BỘ mặt UI khi đã đăng nhập — không chỉ chỉnh homepage nhỏ.",
            "KHÔNG tái sử dụng Navbar hiện tại làm mô hình điều hướng; đề xuất nav mới.",
            "Bản đồ là MỘT tính năng ngang hàng, không phải toàn bộ app shell.",
            "Home logged-in nên là hub; mọi feature chính một click tới (Tìm phòng/map, Chợ đồ, "
            "Tin tìm phòng, Đã lưu, Ở ghép, Tin nhắn, Lịch, Gói Đăng ký + Thông báo, Nhật ký, Hồ sơ, "
            "Tin của tôi/Đăng tin, Admin).",
            "Giữ route cũ hoạt động; không invent feature mới (không nav 'So sánh').",
            "Chỉ UI structure/navigation/layout — không đổi API/auth/map data/business rules.",
            "Guest landing: quyết định riêng (giữ nguyên hay không) — user phải approve.",
            "Chưa implement UI cho đến khi user duyệt plan.",
        ],
    )

    add_heading(doc, "Phụ lục — File neo nhanh", 1)
    add_bullets(
        doc,
        [
            "exe-homeji-web/src/App.tsx — routes",
            "exe-homeji-web/src/lib/mapDeepLinks.ts — section redirects",
            "exe-homeji-web/src/components/layout/AppLayout.tsx, Navbar.tsx, MobileTabBar.tsx",
            "exe-homeji-web/src/pages/HomePage.tsx",
            "exe-homeji-web/src/components/map/AuthenticatedHomeMapShell.tsx, MapOmnibox.tsx, MapAppPanel.tsx",
            "exe-homeji-web/src/api/client.ts, index.ts, types.ts, authSession.ts",
            "exe-homeji-web/vite.config.ts — proxy",
            "exe-homeji-backend/Homeji_BE/src/Homeji.Api/Controllers/*",
            "exe-homeji-backend/Homeji_BE/src/Homeji.Api/Program.cs",
            "PRODUCT.md, DESIGN.md, .cursor/rules/homeji-web-stack-lock.mdc",
        ],
    )

    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
